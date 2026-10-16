"""Word formatting for a Vietnamese thesis chapter.

Conventions (HCMUS / VNU-HCM thesis guideline, 2022): A4; Times New Roman 13;
line spacing 1.5; margins top/bottom/right 2 cm, left 2.5 cm; page number bottom
right; headings numbered 1., 1.1, 1.1.1; table caption ABOVE the table, figure
caption below; IEEE-style references.

User constraints: no icons, no special symbols, no bookmarks. python-docx adds no
bookmarks, and no TOC field is inserted (a TOC field creates hidden _Toc bookmarks
when Word updates it).
"""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"


def _font(run, size=13, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def _para_fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, before=0, after=6, spacing=1.5):
    pf = p.paragraph_format
    pf.alignment = align
    pf.first_line_indent = Cm(1.0) if indent else Cm(0)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing


class Thesis:
    def __init__(self):
        self.doc = Document()
        s = self.doc.sections[0]
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.top_margin = s.bottom_margin = s.right_margin = Cm(2.0)
        s.left_margin = Cm(2.5)
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT; normal.font.size = Pt(13)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        self._page_number_footer(s)
        self.chapter = 0; self.section = 0; self.sub = 0; self.table_no = 0

    def _page_number_footer(self, section):
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
            run._r.append(el)
        _font(run, 13)

    # -- headings ---------------------------------------------------------
    def title(self, text):
        p = self.doc.add_paragraph(); _para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, False, 0, 18, 1.3)
        _font(p.add_run(text.upper()), 16, bold=True)

    def chapter_heading(self, text, new_page=True):
        if new_page and self.chapter > 0:
            self.doc.add_page_break()
        self.chapter += 1; self.section = 0; self.table_no = 0
        p = self.doc.add_paragraph(); _para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, False, 6, 18, 1.3)
        _font(p.add_run(f"CHƯƠNG {self.chapter}. {text.upper()}"), 14, bold=True)

    def section_heading(self, text):
        self.section += 1; self.sub = 0
        p = self.doc.add_paragraph(); _para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT, False, 12, 6)
        p.paragraph_format.keep_with_next = True
        _font(p.add_run(f"{self.chapter}.{self.section}. {text}"), 13, bold=True)

    def sub_heading(self, text):
        self.sub += 1
        p = self.doc.add_paragraph(); _para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT, False, 8, 4)
        p.paragraph_format.keep_with_next = True
        _font(p.add_run(f"{self.chapter}.{self.section}.{self.sub}. {text}"), 13, bold=True, italic=True)

    def unnumbered_heading(self, text):
        p = self.doc.add_paragraph(); _para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, False, 6, 12, 1.3)
        _font(p.add_run(text.upper()), 14, bold=True)

    # -- body -------------------------------------------------------------
    def para(self, text, indent=True):
        p = self.doc.add_paragraph(); _para_fmt(p, indent=indent)
        _font(p.add_run(text), 13)
        return p

    def note(self, text):
        p = self.doc.add_paragraph(); _para_fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, False, 2, 10, 1.15)
        _font(p.add_run(text), 11, italic=True)

    def table(self, caption, header, rows, widths=None, font_size=10.5, note=None):
        self.table_no += 1
        cap = self.doc.add_paragraph(); _para_fmt(cap, WD_ALIGN_PARAGRAPH.CENTER, False, 10, 4, 1.15)
        cap.paragraph_format.keep_with_next = True
        _font(cap.add_run(f"Bảng {self.chapter}.{self.table_no}. "), 12, bold=True)
        _font(cap.add_run(caption), 12, italic=True)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        header_size = min(font_size, 8.5) if len(header) >= 6 else font_size
        for j, h in enumerate(header):
            c = t.rows[0].cells[j]; c.text = ""
            pp = c.paragraphs[0]; _para_fmt(pp, WD_ALIGN_PARAGRAPH.CENTER, False, 1, 1, 1.0)
            _font(pp.add_run(str(h)), header_size, bold=True)
            self._shade(c)
        # Rows never split across a page, and the header row repeats on each page the
        # table spans -- the usual requirement for tables in a printed thesis.
        for i, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit"); trpr.append(cs)
            if i == 0:
                hd = OxmlElement("w:tblHeader"); trpr.append(hd)
        for i, row in enumerate(rows, start=1):
            for j, v in enumerate(row):
                c = t.rows[i].cells[j]; c.text = ""
                pp = c.paragraphs[0]
                _para_fmt(pp, WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER, False, 1, 1, 1.0)
                # (text, True) renders bold -- used instead of an asterisk to mark a
                # result whose confidence interval excludes zero.
                text, bold = (v if isinstance(v, tuple) else (v, False))
                _font(pp.add_run(str(text)), font_size, bold=bold)
        if widths:
            # Word honours cell widths; LibreOffice and some Word versions honour only a
            # FIXED layout with the grid columns set. Do all three or columns come out equal.
            t.autofit = False
            tblpr = t._tbl.tblPr
            layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblpr.append(layout)
            grid = t._tbl.tblGrid
            for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
                if j < len(widths):
                    gc.set(qn("w:w"), str(int(widths[j] * 567)))   # 1 cm = 567 twips
            for j, col in enumerate(t.columns):
                if j < len(widths):
                    col.width = Cm(widths[j])
            for row in t.rows:
                for j, w in enumerate(widths):
                    row.cells[j].width = Cm(w)
        if note:
            self.note(note)
        else:
            sp = self.doc.add_paragraph(); _para_fmt(sp, indent=False, after=4)
        return f"Bảng {self.chapter}.{self.table_no}"

    @staticmethod
    def _shade(cell):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "E7E6E6"); tcpr.append(shd)

    def references(self, refs):
        self.doc.add_page_break()
        self.unnumbered_heading("Tài liệu tham khảo")
        for i, r in enumerate(refs, start=1):
            p = self.doc.add_paragraph(); _para_fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, False, 0, 4, 1.3)
            p.paragraph_format.left_indent = Cm(1.0); p.paragraph_format.first_line_indent = Cm(-1.0)
            _font(p.add_run(f"[{i}] {r}"), 12)

    def save(self, path):
        self.doc.save(path)


# -- number formatting: Vietnamese decimal comma, no special symbols --------
def f(x, d=4):
    if x is None:
        return "không có"
    return f"{x:.{d}f}".replace(".", ",")

def sf(x, d=4):
    if x is None:
        return "không có"
    s = f"{x:+.{d}f}".replace(".", ",")
    return s

def ci(lo, hi, d=4):
    return f"[{sf(lo, d)}; {sf(hi, d)}]"

def ci_plain(lo, hi, d=4):
    return f"[{f(lo, d)}; {f(hi, d)}]"

def pct(x, d=1):
    return f"{100 * x:.{d}f}%".replace(".", ",")

def n(x):
    return f"{int(x):,}".replace(",", ".")
